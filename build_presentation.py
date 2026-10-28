import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Palette Section 9 Cyber-Executive
    COLOR_BG = RGBColor(11, 15, 25)          # Deep Navy Slate
    COLOR_CARD = RGBColor(22, 30, 46)        # Card Slate
    COLOR_CARD_BORDER = RGBColor(38, 52, 78) # Card Border
    COLOR_ACCENT = RGBColor(6, 182, 212)     # Electric Cyan (Section 9)
    COLOR_ACCENT_ALT = RGBColor(56, 189, 248)# Sky Blue
    COLOR_EMERALD = RGBColor(16, 185, 129)   # Mint / Emerald
    COLOR_AMBER = RGBColor(245, 158, 11)     # Amber warning
    COLOR_TEXT_PRIMARY = RGBColor(248, 250, 252)
    COLOR_TEXT_MUTED = RGBColor(148, 163, 184)
    COLOR_TEXT_DIM = RGBColor(100, 116, 139)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="AMARAKI // SECTION 9 ARCHITECTURE"):
        # Category pill/subtitle
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ACCENT

        # Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.7))
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_PRIMARY

    def add_card(slide, left, top, width, height, title="", title_color=COLOR_ACCENT, bg_color=COLOR_CARD, border_color=COLOR_CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)

        if title:
            tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), Inches(0.45))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(16)
            p.font.bold = True
            p.font.color.rgb = title_color
        return card

    # =========================================================================
    # SLIDE 1: Cover
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Decorative accent bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.12), Inches(3.8))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_ACCENT
    bar.line.fill.background()

    tb = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11), Inches(3.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "SECTION 9 // CLUSTER JO3 AI PLATFORM"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT

    p1 = tf.add_paragraph()
    p1.text = "AMARAKI"
    p1.font.size = Pt(56)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_PRIMARY

    p2 = tf.add_paragraph()
    p2.text = "Chief Section 9 — Orchestrateur d'Agents IA & Provisionneur K8s pour Slack"
    p2.font.size = Pt(22)
    p2.font.color.rgb = COLOR_ACCENT_ALT

    p3 = tf.add_paragraph()
    p3.space_before = Pt(24)
    p3.text = "Revue d'Architecture • Écosystème OpenCode • RAG Mnemosyne • Mémoire Git Wiki"
    p3.font.size = Pt(15)
    p3.font.color.rgb = COLOR_TEXT_MUTED

    p4 = tf.add_paragraph()
    p4.space_before = Pt(8)
    p4.text = "Joseph Zacharie — Cluster jo3 (2026)"
    p4.font.size = Pt(13)
    p4.font.color.rgb = COLOR_TEXT_DIM

    # =========================================================================
    # SLIDE 2: Vision & Genèse
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Vision & Genèse : Le Chef d'Orchestre Section 9")

    add_card(s2, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2), "Origine & Philosophie", COLOR_ACCENT)
    tb2_left = s2.shapes.add_textbox(Inches(1.05), Inches(2.4), Inches(5.1), Inches(4.3))
    tf2_l = tb2_left.text_frame
    tf2_l.word_wrap = True
    points_2l = [
        ("Inspiration Ghost in the Shell", "Inspiré du Chef Daisuke Aramaki, stratège impassible de la Section 9, qui coordonne des spécialistes autonomes sur le terrain."),
        ("Le Défi des Agents Décentralisés", "Les assistants IA conversationnels traditionnels manquent d'ancrage opérationnel dans l'infrastructure et risquent de dériver."),
        ("Souveraineté & Auto-Hébergement", "Exécution 100% sur le cluster jo3, sans dépendance opaque externe pour la logique de commande, avec modèles locaux Llama-CPP / LiteLLM."),
        ("Interactivité Naturelle", "Une interface unique : Slack (texte et messages vocaux), accessible sur mobile et desktop à tout moment.")
    ]
    for title, desc in points_2l:
        p_t = tf2_l.add_paragraph()
        p_t.space_before = Pt(10)
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(14)
        p_t.font.color.rgb = COLOR_TEXT_PRIMARY
        p_d = tf2_l.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    add_card(s2, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2), "Objectifs & Principes Clés", COLOR_EMERALD)
    tb2_right = s2.shapes.add_textbox(Inches(7.05), Inches(2.4), Inches(5.2), Inches(4.3))
    tf2_r = tb2_right.text_frame
    tf2_r.word_wrap = True
    points_2r = [
        ("Isolation Stricte (Zero Pollution)", "Chaque tâche génère un Job K8s éphémère. Aucun état résiduel, aucune fuite mémoire entre deux missions."),
        ("Validation Humaine (Human-in-the-Loop)", "Le Chef confirme toujours l'action avant le déploiement sur le cluster, évitant les effets de bord indésirables."),
        ("Spécialisation des Pouvoirs", "Chaque agent possède ses propres serveurs MCP, son outillage CLI (Python, Node, uv) et ses droits minimaux."),
        ("Capitalisation des Connaissances", "Mémoire persistante partagée et versionnée sous forme de Wiki Git.")
    ]
    for title, desc in points_2r:
        p_t = tf2_r.add_paragraph()
        p_t.space_before = Pt(10)
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(14)
        p_t.font.color.rgb = COLOR_TEXT_PRIMARY
        p_d = tf2_r.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 3: Architecture Globale
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Architecture Système & Flux d'Exécution")

    steps = [
        ("1. Slack / Vocal", "Message ou audio reçu\nvia Slack Events API\nsur Traefik Ingress"),
        ("2. Gatekeeper", "Détection d'intention,\ntranscription Whisper\n& machine à états"),
        ("3. Validation", "Confirmation humaine\n(HitL) avant lancement\ndans le thread Slack"),
        ("4. K8s Provisioner", "Création dynamique d'un\nJob Kubernetes éphémère\navec ConfigMaps & Env"),
        ("5. OpenCode Exec", "Exécution du prompt,\nserveurs MCP, skills &\nmémoire Git Wiki"),
        ("6. Reporting", "Reconciliation K8s,\nrestitution du résultat\ndans Slack & télémétrie")
    ]

    card_w = Inches(1.8)
    card_gap = Inches(0.18)
    start_x = Inches(0.8)

    for i, (stitle, sdesc) in enumerate(steps):
        x = start_x + i * (card_w + card_gap)
        add_card(s3, x, Inches(1.8), card_w, Inches(4.0), stitle, COLOR_ACCENT if i % 2 == 0 else COLOR_ACCENT_ALT)
        tb_s = s3.shapes.add_textbox(x + Inches(0.12), Inches(2.6), card_w - Inches(0.24), Inches(3.0))
        tf_s = tb_s.text_frame
        tf_s.word_wrap = True
        p = tf_s.paragraphs[0]
        p.text = sdesc
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_MUTED

    # Bottom summary box
    add_card(s3, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.95), "", bg_color=RGBColor(16, 24, 39))
    tb_bot = s3.shapes.add_textbox(Inches(1.0), Inches(6.1), Inches(11.3), Inches(0.75))
    tf_b = tb_bot.text_frame
    tf_b.word_wrap = True
    p_b = tf_b.paragraphs[0]
    p_b.text = "⚡ Stack Technologique : Rust 2021 (Axum 0.7 + Tokio) • Kube-rs 0.90 • OpenCode CLI • Qdrant / TEI • Tracing OpenTelemetry"
    p_b.font.size = Pt(13)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_EMERALD

    # =========================================================================
    # SLIDE 4: Le Cerveau Gatekeeper
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Le Cerveau Gatekeeper : Analyse, Vocal & Validation")

    add_card(s4, Inches(0.8), Inches(1.7), Inches(3.6), Inches(5.2), "Transcription Vocale", COLOR_ACCENT)
    tb_gk1 = s4.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(3.2), Inches(4.3))
    tf_gk1 = tb_gk1.text_frame
    tf_gk1.word_wrap = True
    tf_gk1.paragraphs[0].text = "• Audio-to-Instruction"
    tf_gk1.paragraphs[0].font.bold = True
    tf_gk1.paragraphs[0].font.size = Pt(14)
    tf_gk1.paragraphs[0].font.color.rgb = COLOR_TEXT_PRIMARY
    p = tf_gk1.add_paragraph()
    p.text = "Interception des messages audio/vidéo Slack dans les canaux désignés."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED

    p = tf_gk1.add_paragraph()
    p.space_before = Pt(10)
    p.text = "• Pipeline Whisper ASR"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_TEXT_PRIMARY
    p = tf_gk1.add_paragraph()
    p.text = "Téléchargement sécurisé avec jeton bot, conversion multipart et appel au endpoint compatible Whisper."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED

    p = tf_gk1.add_paragraph()
    p.space_before = Pt(10)
    p.text = "• Zéro Frottement"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_TEXT_PRIMARY
    p = tf_gk1.add_paragraph()
    p.text = "Permet de piloter les agents en situation de mobilité par de simples mémos vocaux."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED

    add_card(s4, Inches(4.8), Inches(1.7), Inches(3.6), Inches(5.2), "Routage d'Intention", COLOR_ACCENT_ALT)
    tb_gk2 = s4.shapes.add_textbox(Inches(5.0), Inches(2.4), Inches(3.2), Inches(4.3))
    tf_gk2 = tb_gk2.text_frame
    tf_gk2.word_wrap = True
    tf_gk2.paragraphs[0].text = "• Ciblage Explicite"
    tf_gk2.paragraphs[0].font.bold = True
    tf_gk2.paragraphs[0].font.size = Pt(14)
    tf_gk2.paragraphs[0].font.color.rgb = COLOR_TEXT_PRIMARY
    p = tf_gk2.add_paragraph()
    p.text = "Préfixe direct : 'opencode-leclerc: ...', 'opencode-mail: ...', 'opencode-mnemosyne: ...'."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED

    p = tf_gk2.add_paragraph()
    p.space_before = Pt(10)
    p.text = "• Analyse Sémantique Clés"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_TEXT_PRIMARY
    p = tf_gk2.add_paragraph()
    p.text = "Drapeaux sémantiques contextuels : 'courses' -> Leclerc, 'urgence' -> Mail, 'musique' -> Deezer, 'wiki/doc' -> Mnemosyne."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED

    p = tf_gk2.add_paragraph()
    p.space_before = Pt(10)
    p.text = "• Résumé Synthétique"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_TEXT_PRIMARY
    p = tf_gk2.add_paragraph()
    p.text = "Formulation automatique d'une description synthétique de l'action envisagée."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED

    add_card(s4, Inches(8.8), Inches(1.7), Inches(3.7), Inches(5.2), "Validation Humaine (HitL)", COLOR_EMERALD)
    tb_gk3 = s4.shapes.add_textbox(Inches(9.0), Inches(2.4), Inches(3.3), Inches(4.3))
    tf_gk3 = tb_gk3.text_frame
    tf_gk3.word_wrap = True
    tf_gk3.paragraphs[0].text = "• Machine à États"
    tf_gk3.paragraphs[0].font.bold = True
    tf_gk3.paragraphs[0].font.size = Pt(14)
    tf_gk3.paragraphs[0].font.color.rgb = COLOR_TEXT_PRIMARY
    p = tf_gk3.add_paragraph()
    p.text = "State machine par fil Slack : PendingConfirmation -> Confirmed -> Spawned."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED

    p = tf_gk3.add_paragraph()
    p.space_before = Pt(10)
    p.text = "• Contrôle des Ordres"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_TEXT_PRIMARY
    p = tf_gk3.add_paragraph()
    p.text = "Réponse attendue : 'oui', 'go', 'valide' pour déclencher. 'stop', 'annule' pour abandonner."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED

    p = tf_gk3.add_paragraph()
    p.space_before = Pt(10)
    p.text = "• Canaux VIP Directs"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = COLOR_TEXT_PRIMARY
    p = tf_gk3.add_paragraph()
    p.text = "Possibilité d'auto-confirmation sur des canaux administratifs autorisés (#ai)."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 5: Provisioning Kubernetes & Jobs
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Provisioning K8s Natif : Isolation & Cycle de Vie")

    add_card(s5, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2), "Anatomie d'un Job Amaraki", COLOR_ACCENT)
    tb5_l = s5.shapes.add_textbox(Inches(1.05), Inches(2.4), Inches(5.1), Inches(4.3))
    tf5_l = tb5_l.text_frame
    tf5_l.word_wrap = True
    points_5l = [
        ("Nomenclature Strictement Typée", "Job nommé amaraki-<agent_name>-<uuid4>, traçable de bout en bout."),
        ("ConfigMaps Dynamiques", "Déclarés via Helm, montés dans le Pod (/etc/agent/agent.json) pour paramétrer le modèle, les itérations et les variables."),
        ("runner Image Dédiée", "Image de base ghcr.io/jzacharie/opencode avec runtime complet : Node.js LTS, uv, Python 3, Git, curl."),
        ("Placement Intellectuel (NodeSelector)", "Épinglage explicite sur les nœuds amd64 du cluster jo3 pour garantir les performances."),
        ("Reconciliation Loop Asynchrone", "Reconciler Rust en tâche de fond (toutes les 10s) pour capturer la fin d'exécution et renvoyer les logs.")
    ]
    for title, desc in points_5l:
        p_t = tf5_l.add_paragraph()
        p_t.space_before = Pt(8)
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = COLOR_TEXT_PRIMARY
        p_d = tf5_l.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    add_card(s5, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2), "Sécurité & Sandboxing Conteneur", COLOR_AMBER)
    tb5_r = s5.shapes.add_textbox(Inches(7.05), Inches(2.4), Inches(5.2), Inches(4.3))
    tf5_r = tb5_r.text_frame
    tf5_r.word_wrap = True
    points_5r = [
        ("Exécution Non-Root (UID 1000)", "runAsNonRoot: true, runAsUser: 1000, fsGroup: 1000. Aucun accès root dans le conteneur."),
        ("Privilege Escalation Bloquée", "allowPrivilegeEscalation: false et drop: ALL des capabilities Linux."),
        ("Sanitisation des Variables d'Env", "Masquage automatique au démarrage des variables sensibles (*TOKEN*, *PASSWORD*, *SECRET*, *KEY*)."),
        ("RBAC Restreint", "Amaraki dispose d'un ServiceAccount avec droits limités aux Jobs, Pods/log et ConfigMaps dans son namespace."),
        ("Tolérance aux Pannes", "TTL après fin de Job (ttlSecondsAfterFinished) pour nettoyer automatiquement les pods terminés.")
    ]
    for title, desc in points_5r:
        p_t = tf5_r.add_paragraph()
        p_t.space_before = Pt(8)
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = COLOR_TEXT_PRIMARY
        p_d = tf5_r.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 6: Le Runner OpenCode & Runtime
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Le Runner OpenCode : Orchestration In-Container")

    items6 = [
        ("Génération de Configuration", "Construction automatique de `~/.config/opencode/opencode.json` à partir des ConfigMaps : déclaration du modèle LLM, instructions système, timeouts et connecteurs."),
        ("Support Multi-Providers", "Routage natif transparent vers LiteLLM (litellm/), Strata (strata/) et Llama-CPP (llama-cpp/) avec baseURL et API Keys injectées dynamiquement."),
        ("Serveurs MCP Hybrides", "Support des MCP distants (SSE HTTP sur le cluster K8s : Leclerc, OpenObserve, Buzz) et des serveurs MCP locaux (filesystem, git via uvx/npx)."),
        ("Installation Dynamique de Skills", "Import à la volée des compétences via `skills.sh` (`npx -y skills add <skill>`) avant le lancement du prompt agent."),
        ("Isolation de l'Espace de Travail", "Initialisation automatique d'un workspace Git dédié (`~/workspace`) pour que les outils de fichiers et de versioning opèrent proprement.")
    ]

    for i, (title, desc) in enumerate(items6):
        y = Inches(1.7) + i * Inches(1.02)
        add_card(s6, Inches(0.8), y, Inches(11.7), Inches(0.92), "", bg_color=COLOR_CARD)
        tb_it = s6.shapes.add_textbox(Inches(1.1), y + Inches(0.12), Inches(11.1), Inches(0.7))
        tf_it = tb_it.text_frame
        tf_it.word_wrap = True
        p_t = tf_it.paragraphs[0]
        p_t.text = f"{i+1}. {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(14)
        p_t.font.color.rgb = COLOR_ACCENT
        p_d = tf_it.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 7: L'Écosystème des Agents Section 9
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "L'Écosystème des Agents Spécialisés Section 9")

    agents_data = [
        ("agent-code-reviewer", "Revue de Code & Architecture", "Audit qualité, vérification des principes directeurs jo3, détection de bugs et conformité ADR.", "llama-cpp / qwen3.8-27b", COLOR_ACCENT),
        ("opencode-mail", "Urgences Mails & Messagerie", "Connexion sécurisée IMAP/Gmail, filtrage des urgences <48h, synthèse concise et notifications Buzz.", "llama-cpp / qwen3.8-27b", COLOR_AMBER),
        ("opencode-leclerc", "Courses & Panier Leclerc Drive", "Pilotage Chrome via MCP, consultation du panier, recherche catalogue, ajout et ajustement de commandes.", "llama-cpp / qwen3.8-27b", COLOR_EMERALD),
        ("opencode-deezer", "Musique & Synchronisation USB", "Interface avec M-Pacer, consultation des playlists Deezer, génération des manifestes de copie MP3.", "llama-cpp / qwen3.8-27b", COLOR_ACCENT_ALT),
        ("opencode-openobserve", "Observabilité & SRE jo3", "Fouille des streams de logs, métriques ai_code_metrics, analyse des traces OTLP et diagnostic d'incidents.", "strata / qwen3.8-flash", COLOR_ACCENT),
        ("opencode-mnemosyne", "Archiviste & RAG Documentaire", "Interrogation sémantique du lac documentaire, reranking TEI, synthèse sourcée et alimentation du Wiki.", "llama-cpp / qwen3.8-27b", COLOR_EMERALD)
    ]

    for i, (aname, arole, adesc, amodel, acolor) in enumerate(agents_data):
        col = i % 3
        row = i // 3
        x = Inches(0.8) + col * Inches(3.98)
        y = Inches(1.7) + row * Inches(2.65)
        w = Inches(3.78)
        h = Inches(2.45)

        add_card(s7, x, y, w, h, aname, acolor)
        tb_ag = s7.shapes.add_textbox(x + Inches(0.2), y + Inches(0.65), w - Inches(0.4), Inches(1.65))
        tf_ag = tb_ag.text_frame
        tf_ag.word_wrap = True

        p_r = tf_ag.paragraphs[0]
        p_r.text = arole
        p_r.font.bold = True
        p_r.font.size = Pt(12)
        p_r.font.color.rgb = COLOR_TEXT_PRIMARY

        p_d = tf_ag.add_paragraph()
        p_d.space_before = Pt(4)
        p_d.text = adesc
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

        p_m = tf_ag.add_paragraph()
        p_m.space_before = Pt(6)
        p_m.text = f"⚙️ Modèle : {amodel}"
        p_m.font.size = Pt(10)
        p_m.font.color.rgb = acolor

    # =========================================================================
    # SLIDE 8: Synergie RAG avec Mnemosyne
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Intégration RAG Mnemosyne : Savoir & Documentation")

    add_card(s8, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2), "Le Moteur Mnemosyne sur jo3", COLOR_ACCENT)
    tb8_l = s8.shapes.add_textbox(Inches(1.05), Inches(2.4), Inches(5.1), Inches(4.3))
    tf8_l = tb8_l.text_frame
    tf8_l.word_wrap = True
    points_8l = [
        ("Moteur RAG Hybride Haute Performance", "Écrit en Rust, déployé dans le namespace mnemosyne sur le nœud GPU vm168."),
        ("Recherche Vectorielle & Dense (Qdrant)", "Indexation sémantique avec modèle BAAI/bge-m3 via Text Embeddings Inference (TEI)."),
        ("Filtrage & Reranking de Précision", "Combinaison Reciprocal Rank Fusion (RRF) pour garantir les chunks les plus pertinents."),
        ("API REST Légère & Rapide", "Exposition du point d'entrée POST /api/search consommé directement par les conteneurs Amaraki."),
        ("Indexation Enrichie par LLM", "Les chunks contiennent le titre du document, le contexte global et le chunk spécifique pour un ancrage parfait.")
    ]
    for title, desc in points_8l:
        p_t = tf8_l.add_paragraph()
        p_t.space_before = Pt(8)
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = COLOR_TEXT_PRIMARY
        p_d = tf8_l.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    add_card(s8, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2), "Utilisation par les Agents Amaraki", COLOR_EMERALD)
    tb8_r = s8.shapes.add_textbox(Inches(7.05), Inches(2.4), Inches(5.2), Inches(4.3))
    tf8_r = tb8_r.text_frame
    tf8_r.word_wrap = True
    points_8r = [
        ("Variable MNEMOSYNE_URL Déclarée", "Injectée automatiquement dans tous les Jobs (http://mnemosyne.mnemosyne.svc.cluster.local:8080)."),
        ("Agent Dédié opencode-mnemosyne", "Archiviste en chef capable de fouiller des centaines de documents et de produire une synthèse sourcée."),
        ("Augmentation de l'Agent Code Reviewer", "Permet à l'agent de revue de consulter les spécifications d'architecture et les règles de codage jo3 en direct."),
        ("Citations & Vérifiabilité", "Obligation stricte faite aux agents de citer le document source et le chemin NFS (/data/nfs/...)."),
        ("Zéro Hallucination", "Ancrage factuel sur les données réelles du cluster.")
    ]
    for title, desc in points_8r:
        p_t = tf8_r.add_paragraph()
        p_t.space_before = Pt(8)
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = COLOR_TEXT_PRIMARY
        p_d = tf8_r.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 9: Mémoire Persistante Git Wiki
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Mémoire Persistante en Mode Git Wiki pour Agents")

    add_card(s9, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2), "Pourquoi un Wiki Git pour Agents ?", COLOR_ACCENT)
    tb9_l = s9.shapes.add_textbox(Inches(1.05), Inches(2.4), Inches(5.1), Inches(4.3))
    tf9_l = tb9_l.text_frame
    tf9_l.word_wrap = True
    points_9l = [
        ("Transparence & Audit Humain", "Contrairement aux mémoires vectorielles opaques, les fichiers Markdown sont lisibles et éditables par Joseph dans GitHub."),
        ("Versioning Git Complet", "Chaque commit retrace l'agent, le Job ID K8s et l'horodatage exact de l'apprentissage."),
        ("Partage Inter-Agents", "Une procédure découverte par opencode-openobserve dans un playbook peut être lue immédiatement par agent-code-reviewer."),
        ("Zéro Dépendance Base de Données", "Pas de serveur Redis ou PostgreSQL supplémentaire : Git est le socle de vérité universel.")
    ]
    for title, desc in points_9l:
        p_t = tf9_l.add_paragraph()
        p_t.space_before = Pt(10)
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(13.5)
        p_t.font.color.rgb = COLOR_TEXT_PRIMARY
        p_d = tf9_l.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    add_card(s9, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2), "Arborescence & Automatisation GitOps", COLOR_ACCENT_ALT)
    tb9_r = s9.shapes.add_textbox(Inches(7.05), Inches(2.4), Inches(5.2), Inches(4.3))
    tf9_r = tb9_r.text_frame
    tf9_r.word_wrap = True
    tree_text = """Structure normalisée ~/workspace/wiki :
├── index.md                      # Sommaire & table des matières
├── agents/                       # Mémoires individuelles des agents
│   ├── agent-code-reviewer/memory.md
│   ├── opencode-mail/memory.md
│   ├── opencode-leclerc/memory.md
│   └── opencode-mnemosyne/memory.md
├── playbooks/                    # Procédures et guides de secours
│   ├── observability.md
│   └── k8s-triage.md
└── knowledge/                    # Connaissances globales du cluster
    ├── architecture.md
    └── secrets-management.md

Cycle de Vie Automatisé :
1. Clone / Pull sécurisé au lancement du conteneur
2. Lecture du fichier memory.md par l'agent
3. Enrichissement Markdown pendant la mission
4. git commit & push automatique à la sortie du Job !"""
    p_tr = tf9_r.paragraphs[0]
    p_tr.text = tree_text
    p_tr.font.size = Pt(11)
    p_tr.font.name = "Consolas"
    p_tr.font.color.rgb = COLOR_TEXT_PRIMARY

    # =========================================================================
    # SLIDE 10: Observabilité & FinOps
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Observabilité, Télémétrie & OpenObserve")

    add_card(s10, Inches(0.8), Inches(1.7), Inches(3.6), Inches(5.2), "Métriques Prometheus", COLOR_ACCENT)
    tb10_1 = s10.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(3.2), Inches(4.3))
    tf10_1 = tb10_1.text_frame
    tf10_1.word_wrap = True
    points_10_1 = [
        ("Point d'accès /metrics", "Exposition native au format Prometheus scrapé par le collector o2c du cluster."),
        ("amaraki_agents_spawned_total", "Comptabilisation des lancements par agent, modèle et canal Slack."),
        ("amaraki_agent_duration_seconds", "Histogramme de latence d'exécution complète des agents."),
        ("amaraki_agent_errors_total", "Suivi des codes de sortie non nuls des conteneurs.")
    ]
    for title, desc in points_10_1:
        p_t = tf10_1.add_paragraph()
        p_t.space_before = Pt(8)
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = COLOR_TEXT_PRIMARY
        p_d = tf10_1.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    add_card(s10, Inches(4.8), Inches(1.7), Inches(3.6), Inches(5.2), "Traces OpenTelemetry", COLOR_ACCENT_ALT)
    tb10_2 = s10.shapes.add_textbox(Inches(5.0), Inches(2.4), Inches(3.2), Inches(4.3))
    tf10_2 = tb10_2.text_frame
    tf10_2.word_wrap = True
    points_10_2 = [
        ("Pipeline OTLP gRPC/HTTP", "Traces exportées vers OpenObserve (o2-openobserve.p.zacharie.org)."),
        ("Contexte Distribué", "Propagation du traceparent W3C entre la requête Slack, le webhook Axum et le Pod K8s."),
        ("Spans Granulaires", "Mesure précise du temps de transcription Whisper, d'attente de validation et d'exécution OpenCode."),
        ("Corrélation Logs & Traces", "Navigation directe entre span ID et journaux de conteneur.")
    ]
    for title, desc in points_10_2:
        p_t = tf10_2.add_paragraph()
        p_t.space_before = Pt(8)
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = COLOR_TEXT_PRIMARY
        p_d = tf10_2.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    add_card(s10, Inches(8.8), Inches(1.7), Inches(3.7), Inches(5.2), "Tableaux de Bord OpenObserve", COLOR_EMERALD)
    tb10_3 = s10.shapes.add_textbox(Inches(9.0), Inches(2.4), Inches(3.3), Inches(4.3))
    tf10_3 = tb10_3.text_frame
    tf10_3.word_wrap = True
    points_10_3 = [
        ("Stream antigravity_events", "Journalisation unifiée des événements de cycle de vie des agents."),
        ("Stream ai_code_metrics", "Évaluation continue de la qualité du code produit et taux de réussite des tests."),
        ("FinOps & Tokens", "Suivi précis de la consommation des tokens d'inférence (input/output/cache) par modèle."),
        ("Synchronisation GitOps v8", "Dashboards OpenObserve synchronisés dans jo3/values/observability/.")
    ]
    for title, desc in points_10_3:
        p_t = tf10_3.add_paragraph()
        p_t.space_before = Pt(8)
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = COLOR_TEXT_PRIMARY
        p_d = tf10_3.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 11: Sécurité & Secrets
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Sécurité, Secrets & Gouvernance")

    items11 = [
        ("Gestion des Secrets via HashiCorp Vault", "Aucun secret n'est commité dans Git. Les jetons (Slack, Gmail, GitHub, M-Pacer, Groq) sont scellés dans Vault (chemin secret/ai/amaraki)."),
        ("External Secrets Operator (ESO)", "Synchronisation automatisée vers les Secrets Kubernetes natifs (amaraki-slack-credentials, amaraki-github-credentials, opencode-deezer-secret)."),
        ("Masquage Strict des Journaux", "L'audit d'environnement au lancement du Pod remplace automatiquement toute variable contenant TOKEN, SECRET, KEY ou PASSWORD par [REDACTED/MASKED]."),
        ("Authentification Dashboard & APIs", "Accès à l'interface web amaraki.p.zacharie.org protégé par cookie de session chiffré (amaraki_session) ou en-tête Authorization."),
        ("Principe du Moindre Privilège", "Conteneurs en lecture seule sur les ConfigMaps d'instructions, aucun droit d'écriture hors de ~/workspace, communication réseau interne cluster.")
    ]

    for i, (title, desc) in enumerate(items11):
        y = Inches(1.7) + i * Inches(1.02)
        add_card(s11, Inches(0.8), y, Inches(11.7), Inches(0.92), "", bg_color=COLOR_CARD)
        tb_it = s11.shapes.add_textbox(Inches(1.1), y + Inches(0.12), Inches(11.1), Inches(0.7))
        tf_it = tb_it.text_frame
        tf_it.word_wrap = True
        p_t = tf_it.paragraphs[0]
        p_t.text = f"🔒 {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(14)
        p_t.font.color.rgb = COLOR_ACCENT
        p_d = tf_it.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 12: Déploiement GitOps jo3
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Déploiement GitOps sur le Cluster Kubernetes jo3")

    add_card(s12, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2), "ArgoCD & Topologie Helm", COLOR_ACCENT)
    tb12_l = s12.shapes.add_textbox(Inches(1.05), Inches(2.4), Inches(5.1), Inches(4.3))
    tf12_l = tb12_l.text_frame
    tf12_l.word_wrap = True
    points_12l = [
        ("Application ArgoCD Déclarative", "Déclarée dans jo3/Applications/ai/amaraki.yaml, rattachée au projet ai."),
        ("Sources Multiples ArgoCD", "Helm chart hébergé sur helmscharts, fichiers de valeurs values/ai/amaraki/values.yaml et manifestes k8s annexes dans jo3."),
        ("Synchronisation Automatisée (Self-Heal)", "Toute dérive est immédiatement réalignée par ArgoCD avec Prune et ServerSideApply."),
        ("Ingress Traefik & Cert-Manager", "Exposition sécurisée en HTTPS via amaraki.p.zacharie.org avec certificat Let's Encrypt automatique."),
        ("Supervision Homepage", "Intégration directe dans le portail d'accueil avec ping de santé /health.")
    ]
    for title, desc in points_12l:
        p_t = tf12_l.add_paragraph()
        p_t.space_before = Pt(8)
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = COLOR_TEXT_PRIMARY
        p_d = tf12_l.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    add_card(s12, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2), "Cycle CI/CD & Build Conteneurs", COLOR_EMERALD)
    tb12_r = s12.shapes.add_textbox(Inches(7.05), Inches(2.4), Inches(5.2), Inches(4.3))
    tf12_r = tb12_r.text_frame
    tf12_r.word_wrap = True
    points_12r = [
        ("GitHub Actions Automatisé", "Build multi-stage Dockerfile à chaque push sur la branche main de Amaraki."),
        ("Compilation Rust Optimisée", "Profil release avec LTO (Link-Time Optimization), strip des symboles pour une image binaire légère (<25 Mo)."),
        ("Publication GHCR", "Images publiées sur ghcr.io/jzacharie/amaraki et image runner sur ghcr.io/jzacharie/opencode."),
        ("Tests Unitaires & Linting Intégrés", "Exécution automatique de cargo test et clippy avant chaque packaging de conteneur.")
    ]
    for title, desc in points_12r:
        p_t = tf12_r.add_paragraph()
        p_t.space_before = Pt(10)
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(13.5)
        p_t.font.color.rgb = COLOR_TEXT_PRIMARY
        p_d = tf12_r.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 13: Conclusion & Roadmap
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13)
    add_header(s13, "Conclusion & Roadmap Section 9")

    add_card(s13, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2), "Bilan & Acquis Majeurs", COLOR_EMERALD)
    tb13_l = s13.shapes.add_textbox(Inches(1.05), Inches(2.4), Inches(5.1), Inches(4.3))
    tf13_l = tb13_l.text_frame
    tf13_l.word_wrap = True
    points_13l = [
        ("Pont Opérationnel Robuste", "Slack est devenu un véritable terminal de contrôle intelligent pour l'ensemble du cluster."),
        ("Exécution Déterministe & Isolée", "Zéro collision d'état, nettoyage automatique des jobs éphémères."),
        ("RAG Mnemosyne Branché", "Accès instantané à la base documentaire et aux règles d'architecture."),
        ("Mémoire Permanente Versionnée", "Le Wiki Git pour agents offre une continuité cognitive unique et auditable.")
    ]
    for title, desc in points_13l:
        p_t = tf13_l.add_paragraph()
        p_t.space_before = Pt(10)
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(13.5)
        p_t.font.color.rgb = COLOR_TEXT_PRIMARY
        p_d = tf13_l.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    add_card(s13, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2), "Roadmap & Perspectives 2026", COLOR_ACCENT)
    tb13_r = s13.shapes.add_textbox(Inches(7.05), Inches(2.4), Inches(5.2), Inches(4.3))
    tf13_r = tb13_r.text_frame
    tf13_r.word_wrap = True
    points_13r = [
        ("Orchestration Swarm Multi-Agents", "Permettre à Amaraki de chaîner plusieurs agents sur une même mission (ex: Diagnostic SRE -> Revue Code -> Correction Git)."),
        ("Serveurs MCP Dynamiques à Chaud", "Découverte et activation de serveurs MCP à la volée depuis le catalogue jo3."),
        ("Benchmarking Automatisé de Modèles", "Comparateur FinOps/Vitesse/Qualité entre Llama-CPP, LiteLLM et modèles cloud."),
        ("Interface Web Interactive", "Enrichissement du dashboard web pour visualiser les logs live en WebSocket et naviguer dans le Wiki mémoire.")
    ]
    for title, desc in points_13r:
        p_t = tf13_r.add_paragraph()
        p_t.space_before = Pt(10)
        p_t.text = f"• {title}"
        p_t.font.bold = True
        p_t.font.size = Pt(13.5)
        p_t.font.color.rgb = COLOR_TEXT_PRIMARY
        p_d = tf13_r.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    # Save
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    out = os.path.abspath(r"e:\git\JZacharie\Amaraki\Amaraki_Section9_Architecture_Presentation.pptx")
    create_deck(out)
